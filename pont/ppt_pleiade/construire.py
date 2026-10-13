"""Deck PLÉIADE — refonte selon l'avis DESIGNER n°107 (charte PLÉIADE, piste A). Police Arial (repli d'Archivo)."""
import sys
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt

SORTIE = sys.argv[1]
NB = 11
F = "Arial"
NBSP = " "

# ── jetons (avis n°107) ──
CLAIR, SOMBRE = "F3F0E7", "161719"
T_TITRE, T_CORPS, T_SEC, T_PIED = "14140F", "1C1E21", "3B3E46", "5B5F69"
S_1, S_2, S_3 = "EDEEF2", "C3C8D8", "9AA0B4"
GRAPHITE, CRAIE = "1B1C1F", "EDE3C8"
FILET = "C9C3B3"  # filet discret sur fond clair
FILET_S = "3A3D44"

prs = Presentation()
prs.slide_width, prs.slide_height = Pt(960), Pt(540)
VIDE = prs.slide_layouts[6]


def rgb(h):
    return RGBColor.from_string(h)


def fond(s, h):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(h)


def typo(t):
    t = t.replace("'", "’").replace(" :", NBSP + ":").replace(" ;", NBSP + ";")
    return t.replace("« ", "«" + NBSP).replace(" »", NBSP + "»")


def texte(s, x, y, w, h, contenu, taille=24, couleur=T_CORPS, gras=False, align=PP_ALIGN.LEFT,
          ancre=MSO_ANCHOR.TOP, interligne=1.1, espace_apres=0, nom=None):
    """contenu : str, ou liste de paragraphes ; un paragraphe = str ou liste de (texte, {taille, couleur, gras})."""
    b = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    if nom:
        b.name = nom
    tf = b.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ancre
    paras = contenu if isinstance(contenu, list) else [contenu]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = interligne
        p.space_after = Pt(espace_apres)
        runs = para if isinstance(para, list) else [(para, {})]
        for t, o in runs:
            r = p.add_run()
            r.text = typo(t)
            f = r.font
            f.name = F
            f.size = Pt(o.get("taille", taille))
            f.bold = o.get("gras", gras)
            f.color.rgb = rgb(o.get("couleur", couleur))
    return b


def boite(s, x, y, w, h, remplissage=None, trait=None, epaisseur=1, pointille=False, forme=MSO_SHAPE.RECTANGLE, nom=None):
    r = s.shapes.add_shape(forme, Pt(x), Pt(y), Pt(w), Pt(h))
    if nom:
        r.name = nom
    r.shadow.inherit = False
    if remplissage:
        r.fill.solid()
        r.fill.fore_color.rgb = rgb(remplissage)
    else:
        r.fill.background()
    if trait:
        r.line.color.rgb = rgb(trait)
        r.line.width = Pt(epaisseur)
        if pointille:
            r.line.dash_style = MSO_LINE.DASH
    else:
        r.line.fill.background()
    return r


def fleche(s, x, y, w, h, couleur=T_SEC):
    return boite(s, x, y, w, h, remplissage=couleur, forme=MSO_SHAPE.RIGHT_ARROW)


def ligne(s, x1, y1, x2, y2, couleur, ep=2, pointille=False):
    c = s.shapes.add_connector(1, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    c.line.color.rgb = rgb(couleur)
    c.line.width = Pt(ep)
    if pointille:
        c.line.dash_style = MSO_LINE.DASH
    return c


def titre(s, t, sombre=False):
    texte(s, 48, 40, 864, 80, t, taille=32, gras=True, couleur=S_1 if sombre else T_TITRE,
          ancre=MSO_ANCHOR.TOP, interligne=1.05, nom="Titre")


def pied(s, n, sombre=False):
    c = S_3 if sombre else T_PIED
    texte(s, 48, 505, 600, 16, "PLÉIADE · Exercice · non classifié", taille=12, couleur=c)
    texte(s, 812, 505, 100, 16, f"{n}{NBSP}/{NBSP}{NB}", taille=12, couleur=c, align=PP_ALIGN.RIGHT)


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


def constellation(s, x, y, echelle=1.0, couleur=CRAIE):
    # les 7 Pléiades (disposition approchée), motif de la marque
    pts = [(0, 30, 9), (34, 0, 7), (52, 34, 12), (88, 20, 8), (100, 58, 9), (62, 70, 7), (24, 64, 6)]
    for dx, dy, r in pts:
        boite(s, x + dx * echelle, y + dy * echelle, r * echelle, r * echelle, remplissage=couleur, forme=MSO_SHAPE.OVAL)


def nouvelle(sombre=False):
    s = prs.slides.add_slide(VIDE)
    fond(s, SOMBRE if sombre else CLAIR)
    return s


# ════════ 1. Couverture (sombre) ════════
s = nouvelle(True)
constellation(s, 48, 96, 1.0)
texte(s, 48, 196, 864, 90, "PLÉIADE", taille=72, gras=True, couleur=S_1)
texte(s, 48, 290, 864, 40, "L'espace informationnel d'entraînement du CECPC", taille=28, couleur=S_2)
texte(s, 48, 340, 864, 30, f"D'ORION{NBSP}26 à DELATTRE{NBSP}26 · six mois d'avancées", taille=20, couleur=S_2)
texte(s, 48, 505, 864, 20, "Exercice · non classifié · pays, personnages, médias et villes cités : fictifs", taille=12, couleur=S_3)
notes(s, "PLÉIADE est la plateforme d'entraînement du CECPC. Elle reproduit, pour chaque exercice, un espace "
         "informationnel complet et fermé : réseau social, presse, messagerie, fiches de personnages, planche des "
         "injects et outils de mesure. Cette présentation retrace six mois d'évolution, d'ORION 26 à DELATTRE 26.")

# ════════ 2. Vase clos + 4 piliers ════════
s = nouvelle()
titre(s, ["La lutte informationnelle se joue en ligne :", "PLÉIADE la reproduit en vase clos"])
# périmètre fermé
boite(s, 72, 152, 396, 300, trait=T_CORPS, epaisseur=2, nom="Périmètre fermé")
texte(s, 96, 164, 300, 20, "Réseau fermé", taille=18, gras=True, couleur=T_SEC)
for i, m in enumerate(["Réseau social", "Presse", "Messagerie"]):
    boite(s, 96, 200 + i * 78, 348, 60, remplissage=GRAPHITE)
    texte(s, 116, 200 + i * 78, 316, 60, m, taille=24, gras=True, couleur=CRAIE, ancre=MSO_ANCHOR.MIDDLE)
# porte VPN
boite(s, 48, 288, 48, 36, remplissage=CLAIR, trait=T_CORPS, epaisseur=2)
texte(s, 48, 288, 48, 36, "VPN", taille=18, gras=True, couleur=T_CORPS, align=PP_ALIGN.CENTER, ancre=MSO_ANCHOR.MIDDLE)
# 4 piliers
piliers = [("Réalisme", "des outils et des personnages crédibles"), ("Rapidité", "une zone déployée en minutes"),
           ("Isolation", "les fausses rumeurs restent dans la salle"), ("Reproductibilité", "tout se réutilise d'un exercice à l'autre")]
for i, (t, d) in enumerate(piliers):
    y = 152 + i * 78
    boite(s, 504, y + 4, 12, 12, remplissage=T_CORPS)
    texte(s, 530, y - 4, 382, 70, [[(t, {"taille": 24, "gras": True, "couleur": T_TITRE})], [(d, {"taille": 20, "couleur": T_SEC})]])
pied(s, 2)
notes(s, "Le métier se joue sur des réseaux sociaux, des sites d'information et des messageries. PLÉIADE reproduit cet "
         "écosystème dans un espace fermé, accessible par VPN nominatif. Quatre arguments : réalisme (outils et "
         "personnages crédibles), rapidité (une zone et ses applications se déploient en quelques minutes), isolation "
         "(un exercice fabrique de fausses rumeurs et de faux communiqués : ils doivent rester dans la salle), "
         "reproductibilité (personnages, médias et scénarios se réutilisent d'un exercice à l'autre).")

# ════════ 3. Frise ════════
s = nouvelle()
titre(s, "Six mois, quatre exercices : chaque manque a produit un outil")
mois = ["Avr.", "Mai", "Juin", "Juil.", "Août", "Sept.", "Oct."]
x0, x1, ya = 72, 888, 300
pas = (x1 - x0) / 6
ligne(s, x0, ya, x0 + pas * 4.2, ya, T_SEC, 3, pointille=True)
ligne(s, x0 + pas * 4.2, ya, x1, ya, T_CORPS, 4)
for i, m in enumerate(mois):
    texte(s, x0 + i * pas - 34, ya + 12, 68, 22, m, taille=18, couleur=T_PIED, align=PP_ALIGN.CENTER)
# légendes des deux époques, au-dessus de l'axe, entre les repères
texte(s, x0 + 18, ya - 32, 290, 22, "Pages HTML envoyées en ZIP", taille=18, couleur=T_SEC)
texte(s, x1 - 16 - 260, ya - 32, 260, 22, "Serveur unique, accès VPN", taille=18, gras=True, couleur=T_CORPS, align=PP_ALIGN.RIGHT)
etapes = [(0, f"ORION{NBSP}26", "réseau social + premier COCKPIT", True),
          (1, "AURIGE 2BB", "MELMIL, EHO, presse HTML", False),
          (2.3, "MINOTAURE 26", "impact non mesurable", True),
          (4.2, "Fusion", "naissance de PLÉIADE", False),
          (6, f"DELATTRE{NBSP}26", "avatars IA + mesure", True)]
for pos, nom_, outil, haut in etapes:
    cx = x0 + pos * pas
    # jalons évidés, seul l'aboutissement (DELATTRE 26) est plein
    if pos == 6:
        boite(s, cx - 10, ya - 10, 20, 20, remplissage=T_TITRE, forme=MSO_SHAPE.OVAL)
    else:
        boite(s, cx - 8, ya - 8, 16, 16, remplissage=CLAIR, trait=GRAPHITE, epaisseur=2, forme=MSO_SHAPE.OVAL)
    if haut:
        ligne(s, cx, 222, cx, ya - 10, FILET, 1)
        yb = 150
    else:
        ligne(s, cx, ya + 40, cx, 384, FILET, 1)
        yb = 390
    # les repères des extrémités s'alignent sur la marge au lieu de déborder
    if pos == 0:
        tx, al = 48, PP_ALIGN.LEFT
    elif pos == 6:
        tx, al = 912 - 200, PP_ALIGN.RIGHT
    else:
        tx, al = cx - 100, PP_ALIGN.CENTER
    lw = 200
    if not haut:
        tx, lw = cx - 130, 260
    texte(s, tx, yb, lw, 70, [[(nom_, {"taille": 20, "gras": True, "couleur": T_TITRE})], [(outil, {"taille": 18, "couleur": T_SEC})]],
          align=al)
pied(s, 3)
notes(s, "Avril, ORION 26 : le Mastodon prévu ne convient pas à l'exercice ; création de MASTORION (réseau social "
         "d'exercice) et d'un premier COCKPIT de suivi. Mai, AURIGE 2BB : JEMM n'est pas adapté à l'influence et les "
         "joueurs n'ont pas Internet ; création de MELMIL, de l'EHO et d'une presse en pages HTML, diffusés par mail "
         "en ZIP. Juin-juillet, MINOTAURE 26 : une copie par poste, et le RETEX note que l'impact de l'influence n'est "
         "pas mesurable. Août : fusion de MASTAURIGE et MASTORION, naissance de PLÉIADE (nom officialisé le 11 "
         "septembre). Septembre : serveur au CECPC, accès par VPN. Octobre, DELATTRE 26 : avatars animés par IA, "
         "COCKPIT amélioré pour mesurer, MELMIL intégré au jeu, presse supplémentaire, chat sur le réseau social "
         "(besoin exprimé par la LLE).")

# ════════ 4. Avant / après ════════
s = nouvelle()
titre(s, ["D'un ZIP envoyé par mail à une plateforme unique,", "mise à jour en direct"])
lignes = [("Déploiement", "une copie par poste", "une zone en minutes"),
          ("Personnages", "fiches en double", f"une seule fiche (≈{NBSP}4{NBSP}000)"),
          ("Population", "muette sans animateur", "105 avatars IA qui réagissent"),
          ("Injects", "JEMM recopié", "MELMIL importe JEMM"),
          ("Mesure", "impact non mesuré", "opinion suivie au COCKPIT")]
texte(s, 260, 150, 300, 24, "Avant (AURIGE 2BB, MINOTAURE)", taille=18, gras=True, couleur=T_SEC)
boite(s, 576, 144, 336, 332, remplissage=GRAPHITE)
texte(s, 596, 150, 300, 24, f"DELATTRE{NBSP}26", taille=18, gras=True, couleur=CRAIE)
for i, (k, a, b) in enumerate(lignes):
    y = 190 + i * 56
    texte(s, 48, y, 200, 40, k, taille=20, gras=True, couleur=T_TITRE, ancre=MSO_ANCHOR.MIDDLE)
    texte(s, 260, y, 300, 40, a, taille=20, couleur=T_SEC, ancre=MSO_ANCHOR.MIDDLE)
    texte(s, 596, y, 300, 40, b, taille=20, couleur=CRAIE, ancre=MSO_ANCHOR.MIDDLE)
    if i:
        ligne(s, 48, y - 8, 560, y - 8, FILET, 1)
pied(s, 4)
notes(s, "Avant (AURIGE 2BB et MINOTAURE 26) : outils séparés, pages HTML diffusées par mail en ZIP, une copie par "
         "poste ; fiches de personnages faites deux fois ; le réseau ne réagit que si un animateur écrit ; JEMM "
         "synchronisé à la main vers MELMIL ; l'impact de l'influence n'est pas mesurable. Avec DELATTRE 26 : une zone "
         "et ses applications en quelques minutes sur un serveur unique, mises à jour pendant l'exercice ; une fiche "
         "par avatar dans l'EHO, la même partout, environ 4 000 avatars ; 105 avatars animés par IA ; MELMIL alimenté "
         "par les exports JEMM ; le COCKPIT suit l'activité, les hashtags et l'opinion de la population.")

# ════════ 5. Schéma zone ════════
s = nouvelle()
titre(s, ["Une zone par exercice : les applications à la carte,", "une seule identité partout"])
boite(s, 72, 144, 840, 332, trait=T_CORPS, epaisseur=2, nom="Zone")
texte(s, 88, 152, 400, 22, f"Zone « DELATTRE{NBSP}26 »", taille=18, gras=True, couleur=T_SEC)
boite(s, 48, 290, 48, 40, remplissage=CLAIR, trait=T_CORPS, epaisseur=2)
texte(s, 48, 290, 48, 40, "VPN", taille=18, gras=True, align=PP_ALIGN.CENTER, ancre=MSO_ANCHOR.MIDDLE)
groupes = [("Préparer", ["EHO", "MELMIL", "Animation"]), ("Jouer", ["Réseau social", "Presse", "Messagerie"]),
           ("Observer et faire réagir", ["COCKPIT", "Moteur IA"])]
gx = 112
for g, apps in groupes:
    texte(s, gx, 182, 240, 24, g, taille=18, gras=True, couleur=T_TITRE)
    for j, a in enumerate(apps):
        boite(s, gx, 214 + j * 52, 240, 42, remplissage="FFFFFF", trait=FILET)
        texte(s, gx + 14, 214 + j * 52, 220, 42, a, taille=20, couleur=T_CORPS, ancre=MSO_ANCHOR.MIDDLE)
    gx += 262
boite(s, 112, 384, 764, 52, remplissage=GRAPHITE)
texte(s, 132, 384, 734, 52, "EHO : chaque personnage existe une fois, même identité partout",
      taille=20, couleur=CRAIE, ancre=MSO_ANCHOR.MIDDLE)
pied(s, 5)
notes(s, "Une zone est une bulle étanche par exercice : ses comptes, ses données, ses adresses. On y déploie les "
         "applications dont l'exercice a besoin, choisies dans un catalogue de 11 applications (plateforme SaaS "
         "modulaire). Une seule connexion et une barre « Apps de la zone » pour passer de l'une à l'autre (digital "
         "workspace). Chaque personnage existe une seule fois dans l'EHO et garde la même identité partout : il "
         "publie sur le réseau social, signe dans la presse, écrit dans la messagerie. Hors de ce schéma : LEAC, le "
         "logiciel de contrôle des postes de commandement, est aussi une application PLÉIADE.")

# ════════ 6. Flux ════════
s = nouvelle()
titre(s, "Du plan JEMM à l'écran des joueurs, dans un seul circuit")
etapes = [("JEMM", "l'incident est planifié"), ("MELMIL", "on le prépare"), ("Animation", "on programme"),
          ("À l'heure", "publication automatique"), ("Joueurs", "réseau, presse, messagerie"), ("COCKPIT", "on mesure")]
w, g = 120, 28.8
for i, (t, d) in enumerate(etapes):
    x = 48 + i * (w + g)
    boite(s, x, 172, w, 160, remplissage="FFFFFF", trait=FILET)
    texte(s, x + 10, 184, w - 20, 24, f"{i + 1}", taille=18, gras=True, couleur=T_SEC)
    texte(s, x + 10, 216, w - 16, 110, [[(t, {"taille": 20, "gras": True, "couleur": T_TITRE})], [(d, {"taille": 18, "couleur": T_SEC})]])
    if i < 5:
        fleche(s, x + w + 3.4, 242, 22, 20, T_SEC)
boite(s, 48, 374, 864, 52, remplissage=GRAPHITE)
texte(s, 68, 374, 830, 52, "À chaque étape, les mêmes personnages, venus de l'EHO", taille=20, couleur=CRAIE, ancre=MSO_ANCHOR.MIDDLE)
pied(s, 6)
notes(s, "1. L'incident est planifié dans JEMM puis exporté. 2. Il arrive sur la planche MELMIL, qui importe les "
         "exports JEMM (events, storylines, incidents) : on le prépare et on le valide. 3. Dans la console "
         "d'animation, un scénario programme les publications, rattaché à l'incident. 4. Les publications partent "
         "seules, à l'heure. 5. Elles s'affichent sur le réseau social, la presse, la messagerie. 6. Le COCKPIT montre "
         "les réactions ; MELMIL suit le déroulé et les demandes de produits.")

# ════════ 7. 105 (sombre) ════════
s = nouvelle(True)
titre(s, [f"Nouveauté DELATTRE{NBSP}26 :", "un espace informationnel qui réagit"], sombre=True)
texte(s, 48, 174, 400, 150, "105", taille=120, gras=True, couleur=CRAIE)
texte(s, 48, 324, 400, 70, ["avatars animés par IA,", "sans joueur humain derrière"], taille=24, couleur=S_1)
# barre 15 / 70 / 15 sur 100 citoyens
bx, bw, by = 480, 432, 220
parts = [(0.15, "15 %", "pro-Mercure", S_3), (0.70, "70 %", "neutres", S_2), (0.15, "15 %", "pro-Arnland", S_3)]
x = bx
texte(s, bx, 184, bw, 24, "100 citoyens de l'Arnland", taille=18, gras=True, couleur=S_1)
for f_, pct, lib, c in parts:
    boite(s, x, by, bw * f_ - 2, 48, remplissage=c)
    if x == bx:
        lx, lw, al = bx, 140, PP_ALIGN.LEFT
    elif f_ == 0.15:
        lx, lw, al = bx + bw - 140, 140, PP_ALIGN.RIGHT
    else:
        lx, lw, al = x, bw * f_, PP_ALIGN.CENTER
    texte(s, lx, by + 56, lw, 50, [[(pct.replace(" ", NBSP), {"taille": 20, "gras": True, "couleur": S_1})], [(lib, {"taille": 18, "couleur": S_2})]],
          align=al)
    x += bw * f_
texte(s, bx, 354, bw, 60, "+ 3 avatars de test et 2 relais mercuriens", taille=18, couleur=S_2)
pied(s, 7, True)
notes(s, "Jusqu'ici, le réseau ne réagissait que si un animateur écrivait. Sur DELATTRE 26, 105 avatars sont animés par "
         "IA : 100 citoyens de l'Arnland orientés 15 % pro-Mercure, 15 % pro-Arnland, 70 % neutres, plus 3 avatars de "
         "test et 2 relais mercuriens. Ils ne sont joués par personne (groupe à part). Chacun a une fiche complète : "
         "âge, métier, foyer, horaires, caractère, façon d'écrire. Ils vivent : ils travaillent, dorment, parlent de "
         "leur vie autant que de la crise. Comme dans la vraie vie, la plupart lisent beaucoup et écrivent peu.")

# ════════ 8. Programme / IA ════════
s = nouvelle()
titre(s, "Le programme tient la psychologie ; l'IA ne fait que lire et formuler")
blocs = [(48, 170, "L'IA lit", "le fil d'actualité", "FFFFFF"),
         (262, 300, "Le programme décide", ["humeur · mémoire · relations", "opinions · qui réagit"], GRAPHITE),
         (606, 200, "L'IA formule", "dans la voix du personnage", "FFFFFF")]
for x, w, t, d, c in blocs:
    sombre_b = c == GRAPHITE
    boite(s, x, 180, w, 150, remplissage=c, trait=None if sombre_b else FILET)
    ds = d if isinstance(d, list) else [d]
    texte(s, x + 16, 196, w - 32, 130, [[(t, {"taille": 24, "gras": True, "couleur": CRAIE if sombre_b else T_TITRE})]]
          + [[(dd, {"taille": 18, "couleur": S_2 if sombre_b else T_SEC})] for dd in ds])
fleche(s, 229, 246, 22, 20, T_SEC)
fleche(s, 573, 246, 22, 20, T_SEC)
# barrière de garde-fous puis publication
boite(s, 822, 180, 6, 164, remplissage=T_TITRE)
texte(s, 838, 230, 74, 50, "Publié", taille=18, gras=True, couleur=T_TITRE, ancre=MSO_ANCHOR.MIDDLE)
texte(s, 606, 350, 306, 60, ["Garde-fous : relecture, interdits,", "plafonds, interrupteur"], taille=18, couleur=T_SEC, align=PP_ALIGN.RIGHT)
texte(s, 48, 420, 864, 30, "◐ En essai : rédacteur dédié et 36 nuances d'émotion", taille=18, couleur=T_SEC)
pied(s, 8)
notes(s, "Le modèle d'IA ne décide pas seul. Un programme tient la psychologie de chaque avatar, avec des règles "
         "chiffrées et sourcées : l'humeur (qui revient vers un état de repos propre au caractère), la mémoire "
         "(souvenirs classés par importance, consolidés la nuit), les relations avec chaque compte (familiarité, "
         "confiance, rancune), les opinions sur des axes (la Force alliée, Mercure, les institutions) et leurs "
         "bascules avec une phase de doute, et le choix de qui lit et qui réagit. L'IA sert à lire le fil et à "
         "formuler la réaction. Garde-fous : interrupteur général coupé par défaut, interdits absolus (aucune "
         "insulte, aucune information militaire, ne jamais se dire IA, ne rien inventer), plafonds de messages, "
         "relecture automatique avant publication, alerte de dérive. L'animation peut geler, imposer, arrêter. En "
         "essai sur une partie de la population, avec un groupe témoin : un rédacteur dédié et un dictionnaire de "
         "36 nuances d'émotion.")

# ════════ 9. Courbe d'opinion (illustration) ════════
s = nouvelle()
titre(s, "Pour la première fois, on voit l'effet d'une action de joueur sur l'opinion")
cd = CategoryChartData()
cd.categories = ["8 h", "10 h", "12 h", "14 h", "16 h", "18 h", "20 h"]
cd.add_series("Opinion moyenne de la population IA", (-18, -20, -12, 4, 14, 18, 16))
ch = s.shapes.add_chart(XL_CHART_TYPE.LINE, Pt(48), Pt(200), Pt(560), Pt(276), cd).chart
ch.has_legend = False
ch.has_title = False
ch.font.size = Pt(16)
ch.font.name = F
ch.font.color.rgb = rgb(T_SEC)
ser = ch.series[0]
ser.format.line.color.rgb = rgb(T_TITRE)
ser.format.line.width = Pt(4)
ser.smooth = False
va = ch.value_axis
va.minimum_scale, va.maximum_scale, va.major_unit = -100, 100, 50
va.major_gridlines.format.line.color.rgb = rgb(FILET)
va.format.line.fill.background()
ch.category_axis.format.line.color.rgb = rgb(T_PIED)
ch.category_axis.tick_labels.font.size = Pt(18)
ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
va.tick_labels.font.size = Pt(18)
va.tick_labels.number_format = "+0;−0;0"
va.tick_labels.number_format_is_linked = False
texte(s, 48, 150, 560, 22, "Opinion moyenne de la population IA", taille=18, gras=True, couleur=T_TITRE)
texte(s, 48, 174, 560, 22, "ILLUSTRATION · données fictives", taille=18, gras=True, couleur=T_SEC)
# annotation : inject joueur
boite(s, 300, 215, 2, 223, remplissage=T_SEC)
texte(s, 308, 392, 200, 40, "message des joueurs", taille=18, couleur=T_SEC)
# pôles de l'axe d'opinion
texte(s, 112, 214, 130, 22, "▲ soutient", taille=18, couleur=T_SEC)
texte(s, 112, 410, 130, 22, "▼ rejette", taille=18, couleur=T_SEC)
# ce que suit le Reporting
texte(s, 640, 150, 272, 28, "Le COCKPIT suit", taille=20, gras=True, couleur=T_TITRE)
items = ["l'activité par série de comptes", "les hashtags repris", "l'opinion sur chaque axe", "qui la fait bouger : joueurs, IA ou faits"]
for y, t in zip((194, 262, 306, 350), items):
    boite(s, 640, y + 8, 10, 10, remplissage=T_TITRE)
    texte(s, 662, y, 250, 56, t, taille=20, couleur=T_CORPS)
pied(s, 9)
notes(s, "ILLUSTRATION (courbe fictive). Le RETEX de MINOTAURE 26 relevait que l'impact de l'influence n'était pas "
         "mesurable (« so what ? »). Le Reporting du COCKPIT suit désormais l'activité heure par heure par série de "
         "comptes, les hashtags repris (signe qu'une campagne prend), l'opinion de la population simulée sur chaque "
         "axe, qui la fait bouger (joueurs et animation, avatars IA, faits) et les avatars ébranlés. On peut ainsi "
         "montrer à un joueur l'effet de son action sur une population, et pas seulement compter ses publications. "
         "Les relevés réels de DELATTRE 26 sont en cours d'analyse ; ils seront consolidés au RETEX.")

# ════════ 10. Fermé, local, réparti ════════
s = nouvelle()
titre(s, "Fermé, local, réparti : rien ne sort de la salle")
cols = [(48, "Joueurs et animation", "profil VPN nominatif"), (360, "Serveur du CECPC", "les zones d'exercice"),
        (672, "Nœuds de calcul IA", "PC et Mac")]
for i, (x, t, d) in enumerate(cols):
    c = GRAPHITE if i == 1 else "FFFFFF"
    boite(s, x, 170, 240, 130, remplissage=c, trait=None if i == 1 else FILET)
    texte(s, x + 16, 186, 208, 110, [[(t, {"taille": 22, "gras": True, "couleur": CRAIE if i == 1 else T_TITRE})],
                                     [(d, {"taille": 18, "couleur": S_2 if i == 1 else T_SEC})]])
for x in (288, 600):
    fleche(s, x + 25, 225, 22, 20, T_SEC)
    texte(s, x, 252, 72, 22, "VPN", taille=18, gras=True, couleur=T_SEC, align=PP_ALIGN.CENTER)
lignes10 = ["IA locale, licence libre : aucun service en ligne",
            "Une machine s'arrête : une autre reprend ses avatars",
            "Le VPN ne donne aucun droit"]
for i, t in enumerate(lignes10):
    y = 330 + i * 46
    boite(s, 48, y + 9, 10, 10, remplissage=T_TITRE)
    texte(s, 70, y, 842, 40, t, taille=20, couleur=T_CORPS)
pied(s, 10)
notes(s, "On n'entre que par VPN, avec un profil nominatif révocable individuellement : aucune adresse publique, rien "
         "d'indexable. Être sur le VPN ne donne aucun droit : chaque application vérifie elle-même qui vous êtes et ce "
         "que vous pouvez faire (principe zero trust de l'ANSSI) ; comptes de zone anonymes ; l'origine réelle d'une "
         "publication n'est visible que de l'animation ; doctrine de sécurité fondée sur 11 guides de l'ANSSI. Le "
         "modèle d'IA (Gemma 4, Google, licence libre Apache 2.0) tourne sur des ordinateurs de l'équipe, jamais dans "
         "un service en ligne. Le serveur n'ayant pas de carte graphique, il confie le calcul à des nœuds joints par "
         "le VPN, protégés par un jeton et un pare-feu. Depuis le 13 octobre, deux machines forment un parc : chaque "
         "avatar est rattaché à une machine, repris automatiquement par une autre si elle s'arrête ; les cycles "
         "tournent environ deux fois plus vite.")

# ════════ 11. Conclusion (sombre) ════════
s = nouvelle(True)
titre(s, "Ce que PLÉIADE permet", sombre=True)
verbes = [("Entraîner", "un état-major à la lutte informationnelle"), ("Monter", "un exercice vite, et tout réutiliser"),
          ("Faire vivre", "l'espace avec une population simulée"), ("Mesurer", "l'effet des joueurs sur l'opinion")]
for i, (v, d) in enumerate(verbes):
    y = 150 + i * 58
    texte(s, 48, y, 200, 50, v, taille=28, gras=True, couleur=CRAIE)
    texte(s, 250, y + 7, 380, 50, d, taille=20, couleur=S_1)
boite(s, 660, 150, 252, 200, trait=FILET_S, epaisseur=1)
texte(s, 680, 166, 220, 180, [[("La suite", {"taille": 20, "gras": True, "couleur": S_1})],
                               [("◐ en essai : rédacteur, émotions", {"taille": 18, "couleur": S_2})],
                               [("○ à l'étude : lecture d'images", {"taille": 18, "couleur": S_2})],
                               [("→ ensuite : plus d'avatars", {"taille": 18, "couleur": S_2})]],
      espace_apres=10)
constellation(s, 820, 440, 0.6)
pied(s, 11, True)
notes(s, "PLÉIADE permet d'entraîner un état-major à la lutte informationnelle dans un espace réaliste et fermé ; de "
         "monter un exercice vite et de réutiliser personnages, médias et scénarios ; de faire vivre l'espace avec "
         "une population simulée, sans dépendre du nombre d'animateurs ; de mesurer l'effet des actions des joueurs "
         "sur l'opinion ; et d'évoluer en continu, les améliorations arrivant pendant l'exercice. En essai : le "
         "rédacteur dédié et le dictionnaire des émotions, sur une partie de la population avec un groupe témoin. À "
         "l'étude : la lecture des images publiées par les joueurs. Ensuite : une population plus nombreuse et de "
         "nouvelles machines de calcul.")

prs.save(SORTIE)
print("ok", SORTIE, len(prs.slides), "diapos")
